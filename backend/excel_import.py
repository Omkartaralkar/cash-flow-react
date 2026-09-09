"""
Shared Excel-import parsing logic.

This is the same column-detection / row-parsing logic that shipped in
the original import_swamiraj.py, pulled out into its own module so it
can be reused by:

  * import_swamiraj.py   - the original standalone CLI script (kept
                            working exactly as before)
  * routes_import.py     - a new /api/import/excel endpoint that lets
                            any logged-in user import a workbook for
                            their own account from the Settings page

openpyxl only. No pandas.
"""

from __future__ import annotations

from datetime import date, datetime
from decimal import Decimal, InvalidOperation

from openpyxl import load_workbook


def clean_text(value) -> str:
    if value is None:
        return ""
    return str(value).strip()


def parse_amount(value) -> float:
    if value is None or value == "":
        return 0.0

    if isinstance(value, (int, float)):
        return float(value)

    text = clean_text(value)

    text = (
        text.replace("₹", "")
        .replace(",", "")
        .replace("INR", "")
        .replace("Rs.", "")
        .replace("Rs", "")
        .strip()
    )

    if not text:
        return 0.0

    negative = text.startswith("(") and text.endswith(")")

    if negative:
        text = text[1:-1].strip()

    try:
        amount = float(Decimal(text))
    except InvalidOperation:
        raise ValueError(f"Invalid amount: {value!r}")

    return -amount if negative else amount


def parse_date(value) -> str:
    if value is None or value == "":
        return ""

    if isinstance(value, datetime):
        return value.strftime("%Y-%m-%d")

    if isinstance(value, date):
        return value.strftime("%Y-%m-%d")

    text = clean_text(value)

    formats = [
        "%d/%m/%y",
        "%d/%m/%Y",
        "%d-%m-%y",
        "%d-%m-%Y",
        "%d.%m.%y",
        "%d.%m.%Y",
    ]

    for fmt in formats:
        try:
            return datetime.strptime(text, fmt).strftime("%Y-%m-%d")
        except ValueError:
            pass

    if isinstance(value, (int, float)):
        try:
            from openpyxl.utils.datetime import from_excel

            converted = from_excel(value)

            if isinstance(converted, (datetime, date)):
                return converted.strftime("%Y-%m-%d")
        except Exception:
            pass

    raise ValueError(f"Unsupported date format: {value!r}")


def normalize_header(value) -> str:
    return (
        clean_text(value)
        .lower()
        .replace(" ", "")
        .replace("_", "")
    )


def find_columns(ws):
    aliases = {
        "date": {"date"},
        "reason": {"reason", "कारण", "reason/कारण"},
        "income": {"income", "जमा", "income/जमा"},
        "expense": {"expense", "खर्च", "expense/खर्च"},
    }

    for row_number in range(1, min(ws.max_row, 10) + 1):
        found = {}

        for col in range(1, ws.max_column + 1):
            key = normalize_header(ws.cell(row_number, col).value)

            for name, possible in aliases.items():
                normalized_possible = {
                    normalize_header(item) for item in possible
                }

                if key in normalized_possible:
                    found[name] = col

        if all(name in found for name in aliases):
            return row_number, found

    raise ValueError(
        "Could not find required columns. "
        "Expected Date, Reason / कारण, Income / जमा, "
        "and Expense / खर्च."
    )


def make_transaction(tx_date, reason, amount, tx_type):
    return {
        "date": tx_date,
        "name": "",
        "reason": reason,
        "amount": amount,
        "type": tx_type,
    }


def parse_workbook(file_or_path):
    """
    Parse an .xlsx workbook (path or file-like object) and return
    (transactions, skipped_row_count). Raises ValueError on anything
    the sheet doesn't support.
    """

    workbook = load_workbook(
        file_or_path,
        data_only=True,
        read_only=True,
    )

    try:
        ws = workbook.active

        header_row, columns = find_columns(ws)

        transactions = []
        skipped = 0

        for row_number in range(header_row + 1, ws.max_row + 1):
            raw_date = ws.cell(row_number, columns["date"]).value
            raw_reason = ws.cell(row_number, columns["reason"]).value
            raw_income = ws.cell(row_number, columns["income"]).value
            raw_expense = ws.cell(row_number, columns["expense"]).value

            if all(
                value is None or clean_text(value) == ""
                for value in (raw_date, raw_reason, raw_income, raw_expense)
            ):
                continue

            try:
                tx_date = parse_date(raw_date)
                reason = clean_text(raw_reason)

                income = parse_amount(raw_income)
                expense = parse_amount(raw_expense)

                if income and expense:
                    raise ValueError(
                        "Both Income and Expense contain amounts."
                    )

                if income:
                    transactions.append(
                        make_transaction(tx_date, reason, income, "income")
                    )
                elif expense:
                    transactions.append(
                        make_transaction(tx_date, reason, expense, "expense")
                    )
                else:
                    skipped += 1

            except ValueError as exc:
                raise ValueError(f"Excel row {row_number}: {exc}") from exc

        return transactions, skipped

    finally:
        workbook.close()
